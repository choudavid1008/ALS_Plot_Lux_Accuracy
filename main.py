import sys
import os
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QTabWidget, QTableWidget, QTableWidgetItem,
    QLineEdit, QLabel, QFileDialog, QFormLayout, QGroupBox,
    QScrollArea, QMessageBox, QGridLayout, QColorDialog
)
from PySide6.QtCore import Qt, QEvent
from PySide6.QtGui import QColor
from openpyxl.cell.cell import Cell
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from openpyxl import load_workbook
from pathlib import Path

def get_excel_column_name(n):
    """Convert a 0-indexed column number to an Excel column name (0 -> A, 25 -> Z, 26 -> AA)."""
    name = ""
    n += 1
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        name = chr(65 + remainder) + name
    return name

class SelectableLineEdit(QLineEdit):
    def __init__(self, contents="", parent=None):
        super().__init__(contents, parent)

class ColorLineEdit(QLineEdit):
    def __init__(self, color_hex="#808080", parent=None):
        super().__init__(color_hex, parent)
        self.setReadOnly(True) # Make it read-only to force dialog usage
        self.setCursor(Qt.PointingHandCursor)
        self.update_style(color_hex)

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            color = QColorDialog.getColor(QColor(self.text()), self, "Select Color")
            if color.isValid():
                hex_color = color.name().upper()
                self.setText(hex_color)
                self.update_style(hex_color)
        super().mousePressEvent(event)

    def update_style(self, hex_color):
        # Calculate luminance to decide text color (black or white)
        color = QColor(hex_color)
        lum = 0.299 * color.red() + 0.587 * color.green() + 0.114 * color.blue()
        text_color = "black" if lum > 128 else "white"
        self.setStyleSheet(f"background-color: {hex_color}; color: {text_color}; border: 1px solid gray; font-weight: bold;")

class PlotCanvas(FigureCanvas):
    def __init__(self, parent=None, width=14, height=12, dpi=100):
        self.fig, self.axes = plt.subplots(2, 2, figsize=(width, height), dpi=dpi)
        super(PlotCanvas, self).__init__(self.fig)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Lux Accuracy Tool")
        self.resize(1500, 950)

        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)
        self.main_layout = QVBoxLayout(self.central_widget)

        # Top Bar
        self.top_layout = QHBoxLayout()
        self.load_btn = QPushButton("Load Excel File (讀檔)")
        self.load_btn.clicked.connect(self.load_file)
        self.top_layout.addWidget(self.load_btn)

        self.run_btn = QPushButton("Run & Generate Plot (執行)")
        self.run_btn.clicked.connect(self.run_process)
        self.top_layout.addWidget(self.run_btn)

        self.main_layout.addLayout(self.top_layout)

        # Tab Widget
        self.tabs = QTabWidget()
        self.main_layout.addWidget(self.tabs)

        # Data View Tab
        self.data_tab = QWidget()
        self.data_layout = QHBoxLayout(self.data_tab)
        self.tabs.addTab(self.data_tab, "Data View (表格與設定)")

        # Left Panel: Config
        self.config_scroll = QScrollArea()
        self.config_scroll.setWidgetResizable(True)
        self.config_scroll.setFixedWidth(420)
        self.config_container = QWidget()
        self.config_vbox = QVBoxLayout(self.config_container)
        self.config_scroll.setWidget(self.config_container)

        self.config_group = QGroupBox("Configuration (設定)")
        self.config_layout = QFormLayout(self.config_group)

        # Interactive LineEdits
        self.cct_range_input = SelectableLineEdit("C5:C19")
        self.cr_cell = SelectableLineEdit("C22")
        self.cg_cell = SelectableLineEdit("C23")
        self.cb_cell = SelectableLineEdit("C24")
        self.cc_cell = SelectableLineEdit("C25")
        self.cwb_cell = SelectableLineEdit("C26")
        self.lux_range_input = SelectableLineEdit("I31:I75")
        self.reported_range_input = SelectableLineEdit("U31:U75")
        self.subtitle_input = QLineEdit("ALS Summary_Coef(ARRI + XRite_Low + TPE_MFG + VN_MFG)_Verify(XRite_Low)")

        self.config_layout.addRow("CCT Pattern Range:", self.cct_range_input)
        self.config_layout.addRow("Cr Cell:", self.cr_cell)
        self.config_layout.addRow("Cg Cell:", self.cg_cell)
        self.config_layout.addRow("Cb Cell:", self.cb_cell)
        self.config_layout.addRow("Cc Cell:", self.cc_cell)
        self.config_layout.addRow("Cwb Cell:", self.cwb_cell)
        self.config_layout.addRow("Ref Lux Range:", self.lux_range_input)
        self.config_layout.addRow("Reported Lux Range:", self.reported_range_input)
        self.config_layout.addRow("Plot Subtitle:", self.subtitle_input)

        # Color Settings Group
        self.color_group = QGroupBox("CCT Color Settings (顏色設定)")
        self.color_layout = QGridLayout(self.color_group)
        self.color_layout.addWidget(QLabel("CCT (K)"), 0, 0)
        self.color_layout.addWidget(QLabel("Pick Color (點擊設定)"), 0, 1)

        self.color_inputs = []
        default_colors = [("3000", "#FF0000"), ("4000", "#0066FF"), ("4150", "#00AA00")]
        for i, (cct, color) in enumerate(default_colors):
            cct_in = QLineEdit(cct)
            col_in = ColorLineEdit(color)
            self.color_layout.addWidget(cct_in, i+1, 0)
            self.color_layout.addWidget(col_in, i+1, 1)
            self.color_inputs.append((cct_in, col_in))

        self.add_color_btn = QPushButton("+ Add Color Mapping")
        self.add_color_btn.clicked.connect(self.add_color_row)

        self.config_vbox.addWidget(self.config_group)
        self.config_vbox.addWidget(self.color_group)
        self.config_vbox.addWidget(self.add_color_btn)
        self.config_vbox.addStretch()

        # Track which line edit is focused
        self.focused_line_edit = None
        for le in [self.cct_range_input, self.cr_cell, self.cg_cell, self.cb_cell,
                   self.cc_cell, self.cwb_cell, self.lux_range_input, self.reported_range_input]:
            le.installEventFilter(self)

        self.data_layout.addWidget(self.config_scroll)

        # Right Panel: Table View
        self.table_widget = QTableWidget()
        self.table_widget.itemSelectionChanged.connect(self.on_selection_changed)
        self.data_layout.addWidget(self.table_widget)

        # Plot View Tab
        self.plot_tab = QWidget()
        self.plot_layout = QVBoxLayout(self.plot_tab)
        self.tabs.addTab(self.plot_tab, "Plot View (圖表)")

        self.canvas = PlotCanvas(self.plot_tab, width=14, height=12)
        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setWidget(self.canvas)
        self.plot_layout.addWidget(self.scroll_area)

        # State
        self.wb = None
        self.file_path = None

    def add_color_row(self):
        row = len(self.color_inputs) + 1
        cct_in = QLineEdit()
        col_in = ColorLineEdit("#808080")
        self.color_layout.addWidget(cct_in, row, 0)
        self.color_layout.addWidget(col_in, row, 1)
        self.color_inputs.append((cct_in, col_in))

    def eventFilter(self, obj, event):
        if event.type() == QEvent.FocusIn:
            if isinstance(obj, SelectableLineEdit):
                self.set_focused_line_edit(obj)
        return super().eventFilter(obj, event)

    def set_focused_line_edit(self, le):
        if self.focused_line_edit:
            self.focused_line_edit.setStyleSheet("")
        self.focused_line_edit = le
        if self.focused_line_edit:
            self.focused_line_edit.setStyleSheet("background-color: #e6f3ff; border: 2px solid #0078d7;")

    def on_selection_changed(self):
        if not self.focused_line_edit:
            return
        selected_ranges = self.table_widget.selectedRanges()
        if not selected_ranges:
            return
        sel = selected_ranges[0]
        top = sel.topRow() + 1
        bottom = sel.bottomRow() + 1
        left = sel.leftColumn()
        right = sel.rightColumn()
        col_start = get_excel_column_name(left)
        col_end = get_excel_column_name(right)
        if top == bottom and left == right:
            range_str = f"{col_start}{top}"
        else:
            range_str = f"{col_start}{top}:{col_end}{bottom}"
        self.focused_line_edit.setText(range_str)

    def to_float(self, val):
        try:
            return float(val)
        except Exception:
            return np.nan

    def get_range_values(self, ws, cell_range):
        values = []
        try:
            items = ws[cell_range]
            if isinstance(items, Cell):
                values.append(items.value)
            else:
                for row in items:
                    if isinstance(row, Cell):
                         values.append(row.value)
                    else:
                        for cell in row:
                            values.append(cell.value)
        except Exception as e:
            print(f"Error reading range {cell_range}: {e}")
        return values

    def get_single_cell_value(self, ws, cell_addr):
        try:
            item = ws[cell_addr]
            if isinstance(item, Cell):
                return item.value
            else:
                return item[0][0].value
        except Exception as e:
            print(f"Error reading cell {cell_addr}: {e}")
            return None

    def build_cct_cycle(self, ws, data_len):
        pattern_range = self.cct_range_input.text()
        pattern_vals = self.get_range_values(ws, pattern_range)
        pattern = [
            int(self.to_float(v))
            for v in pattern_vals
            if not pd.isna(self.to_float(v))
        ]
        if not pattern:
            return [], [], 0

        pattern_len = len(pattern)
        cct_list = [pattern[i % len(pattern)] for i in range(data_len)]

        color_map = {}
        for cct_in, col_in in self.color_inputs:
            try:
                cct_val = int(cct_in.text())
                color_map[cct_val] = col_in.text()
            except ValueError:
                continue

        color_list = [color_map.get(cct, "#808080") for cct in cct_list]
        return cct_list, color_list, pattern_len

    def extract_sheet_data(self, ws):
        coeff_dict = {
            "Cr": self.to_float(self.get_single_cell_value(ws, self.cr_cell.text())),
            "Cg": self.to_float(self.get_single_cell_value(ws, self.cg_cell.text())),
            "Cb": self.to_float(self.get_single_cell_value(ws, self.cb_cell.text())),
            "Cc": self.to_float(self.get_single_cell_value(ws, self.cc_cell.text())),
            "Cwb": self.to_float(self.get_single_cell_value(ws, self.cwb_cell.text())),
        }

        lux_values = self.get_range_values(ws, self.lux_range_input.text())
        reported_values = self.get_range_values(ws, self.reported_range_input.text())

        df_tmp = pd.DataFrame({
            "lux": lux_values,
            "reported": reported_values
        }).dropna()

        lux_values = df_tmp["lux"].tolist()
        reported_values = df_tmp["reported"].tolist()
        data_len = len(lux_values)

        cct_list, color_list, pattern_len = self.build_cct_cycle(ws, data_len)
        if not cct_list:
             return coeff_dict, pd.DataFrame()

        df = pd.DataFrame({
            "CL-200A Lux": lux_values,
            "Reported_LUX": reported_values,
            "CCT": cct_list,
            "Color": color_list
        })
        df["CL-200A Lux"] = pd.to_numeric(df["CL-200A Lux"], errors="coerce")
        df["Reported_LUX"] = pd.to_numeric(df["Reported_LUX"], errors="coerce")
        df = df.dropna()
        return coeff_dict, df

    def plot_lux_accuracy_on_ax(self, ax, df, sheet_name, file_name, coeff_dict):
        if df.empty:
            ax.set_title(f"{sheet_name}\nNo data")
            ax.grid(True)
            return

        x = df["CL-200A Lux"]
        y = df["Reported_LUX"]
        max_val = max(x.max(), y.max()) * 1.1 if not df.empty else 100
        x_line = np.linspace(0, max_val, 200)

        for cct, group in df.groupby("CCT"):
            ax.scatter(group["CL-200A Lux"], group["Reported_LUX"],
                       color=group["Color"].iloc[0], s=5, alpha=0.85, label=f"{cct}K")

        ax.plot(x_line, x_line, "k--", linewidth=1.2, label="Ideal")
        ax.fill_between(x_line, 0.9 * x_line, 1.1 * x_line, color="gray", alpha=0.2, label="±10%")
        ax.set_xlabel("Reference Lux (CL-200A)")
        ax.set_ylabel("Reported Lux")
        ax.set_title(f"{sheet_name}\n{file_name}", fontsize=10)
        ax.grid(True, alpha=0.3)
        ax.legend(fontsize=8, loc="lower right")

        coeff_text = (
            f"Cr: {coeff_dict['Cr']:.3f}\n"
            f"Cg: {coeff_dict['Cg']:.3f}\n"
            f"Cb: {coeff_dict['Cb']:.3f}\n"
            f"Cc: {coeff_dict['Cc']:.3f}\n"
            f"Cwb: {coeff_dict['Cwb']:.3f}"
        )
        ax.text(0.03, 0.95, coeff_text, transform=ax.transAxes, fontsize=8,
                bbox=dict(boxstyle="round", facecolor="white", alpha=0.85), verticalalignment="top")

    def run_process(self):
        if not self.wb:
            QMessageBox.warning(self, "Warning", "Please load an Excel file first.")
            return

        sheets = ["Negroni (red)", "Pine (green)", "Haze midnight (black)", "Silver"]
        file_name = os.path.basename(self.file_path)

        self.canvas.fig.clf()
        self.canvas.axes = self.canvas.fig.subplots(2, 2)
        axes = self.canvas.axes.flatten()

        for i, sheet in enumerate(sheets):
            ax = axes[i]
            if sheet not in self.wb.sheetnames:
                ax.set_title(f"{sheet}\nNot found")
                ax.grid(True)
                continue

            ws = self.wb[sheet]
            coeff_dict, df = self.extract_sheet_data(ws)
            self.plot_lux_accuracy_on_ax(ax, df, sheet, file_name, coeff_dict)

        self.canvas.fig.suptitle(self.subtitle_input.text(), fontsize=16)
        self.canvas.fig.tight_layout(rect=[0, 0, 1, 0.96])
        self.canvas.draw()

        # Save file
        os.makedirs("output_file", exist_ok=True)
        out_path = os.path.join("output_file", "lux_accuracy_for_all.png")
        self.canvas.fig.savefig(out_path, dpi=300, bbox_inches="tight")

        # Switch tab
        self.tabs.setCurrentIndex(1)
        QMessageBox.information(self, "Success", f"Plot generated and saved to {out_path}")

    def load_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Open Excel File", "", "Excel Files (*.xlsx *.xls)")
        if file_path:
            self.file_path = file_path
            try:
                self.wb = load_workbook(file_path, data_only=True)
                self.display_sheet_in_table(self.wb.active)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Could not load file: {e}")

    def display_sheet_in_table(self, ws):
        self.table_widget.blockSignals(True)
        self.table_widget.setRowCount(ws.max_row)
        self.table_widget.setColumnCount(ws.max_column)

        col_headers = [get_excel_column_name(i) for i in range(ws.max_column)]
        self.table_widget.setHorizontalHeaderLabels(col_headers)
        row_headers = [str(i + 1) for i in range(ws.max_row)]
        self.table_widget.setVerticalHeaderLabels(row_headers)

        for i, row in enumerate(ws.iter_rows()):
            for j, cell in enumerate(row):
                val = str(cell.value) if cell.value is not None else ""
                self.table_widget.setItem(i, j, QTableWidgetItem(val))
        self.table_widget.blockSignals(False)

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
